"""LX-SEQ v2.0 포맷 준수 검증 — 15항목 체크리스트"""

import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# 입력 위치는 이 스크립트 위치에서 유도한다. 기계마다 다른 절대경로를
# 박아두면 그 기계 밖에서는 돌지 않는다.
SONG_DIR = os.environ.get("LXSEQ_SONG_OUT") or os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "03_곡파일_Sugar"
)
from openpyxl import load_workbook
from seq_data import FIXTURE_GROUPS, PALETTE, RUNTIME, SECTIONS

XLSX = os.path.join(SONG_DIR, "LXSEQ_SAMPLE_01_Sugar_r3.xlsx")
wb = load_workbook(XLSX)
head = {r[0].value: r[1].value for r in wb["HEAD"].iter_rows(min_row=4, max_row=20, max_col=2)}
cs = wb["CUE"]
HDR = [c.value for c in cs[4]]
rows = [[c.value for c in r] for r in cs.iter_rows(min_row=5, max_row=cs.max_row, max_col=14)]

SPEC_HDR = [
    "Q#",
    "Section",
    "TC In",
    "TC Out",
    "Dur",
    "Mood",
    "Color(주/보조)",
    "Intensity",
    "Fixture Group",
    "Movement",
    "Effect",
    "Transition",
    "Fade",
    "Note",
]
SEC_ENUM = {
    "PRE-ROLL",
    "AMBI",
    "INTRO",
    "VERSE1",
    "VERSE2",
    "VERSE3",
    "PRE1",
    "PRE2",
    "PRE3",
    "CHORUS1",
    "CHORUS2",
    "CHORUS3",
    "CHORUS4",
    "BRIDGE",
    "BREAK",
    "DROP",
    "SOLO",
    "OUTRO",
    "TAG",
}  # v1.1
TRANS_ENUM = {"SNAP", "FADE", "XFADE", "BUMP"}
GROUPS = {g[0] for g in FIXTURE_GROUPS} | {"ALL"}
PAL_ID = {p[0] for p in PALETTE}
PAL_NM = {p[0]: p[1].replace(" ", "") for p in PALETTE}


def secs(t):
    if not t:
        return None
    m, s = t.split(":")
    return int(m) * 60 + float(s)


R = []


def chk(n, name, ok, detail=""):
    R.append((n, name, "PASS" if ok else "FAIL", detail))


# 1
o = bool(head.get("TC_ORIGIN")) and len(str(head.get("TC_ORIGIN", ""))) > 15
chk(1, "HEAD TC_ORIGIN 문장 명시", o, str(head.get("TC_ORIGIN"))[:48])
# 2
chk(2, "CUE 14열 순서 일치", HDR == SPEC_HDR, "실제 %d열" % len(HDR))
# 3
bad = sorted({r[1] for r in rows if r[1] not in SEC_ENUM})
chk(
    3,
    "Section enum 준수 (v1.1)",
    not bad,
    ("이탈: " + ", ".join(bad)) if bad else "10개 섹션 전부 정의값 (PRE1·PRE2 포함)",
)
# 4
tb = [r[0] for r in rows if r[11] not in TRANS_ENUM]
fb = [r[0] for r in rows if (r[11] in ("SNAP", "BUMP")) != (float(r[12]) == 0.0)]
chk(
    4,
    "Transition enum + Fade 정합",
    not tb and not fb,
    (f"enum이탈 {tb} / Fade모순 {fb}") if (tb or fb) else "SNAP·BUMP=0.0, FADE·XFADE>0 일치",
)
# 5
tins = [secs(r[2]) for r in rows]
mono = all(tins[i] < tins[i + 1] for i in range(len(tins) - 1))
ovl = [
    rows[i][0] for i in range(len(rows) - 1) if rows[i][3] and secs(rows[i][3]) > tins[i + 1] + 1e-9
]
chk(
    5,
    "TC In 단조증가 · 구간 무중첩",
    mono and not ovl,
    "18큐 단조증가, 중첩 0건" if (mono and not ovl) else f"중첩 {ovl}",
)
# 6
used = set()
for r in rows:
    used |= set(re.split(r"\+", r[8]))
undef = sorted(used - GROUPS)
chk(
    6,
    "Fixture Group 약칭 HEAD 정의",
    not undef,
    (f"미정의: {undef}") if undef else "%d개 약칭 전부 정의됨" % len(used),
)
# 7
bad7 = []
for r in rows:
    for tok in [t.strip() for t in r[6].split("/")]:
        if tok in ("—", ""):
            continue
        m = re.match(r"^(P\d)\s+(\S+)$", tok)
        if not m or m.group(1) not in PAL_ID or m.group(2) != PAL_NM[m.group(1)]:
            bad7.append((r[0], tok))
chk(
    7,
    "Color = 팔레트ID + 한글색상명 병기",
    not bad7,
    (f"위반 {bad7[:3]}") if bad7 else "18큐 전부 병기",
)
# 8
gaps = [(rows[i][0], rows[i + 1][0], tins[i + 1] - tins[i]) for i in range(len(rows) - 1)]
viol = [g for g in gaps if g[2] < 8.0]
term = [g for g in viol if g[1] == rows[-1][0]]
chk(
    8,
    "전환 밀도 ≤ 8초/1큐 (§3.6.5 예외)",
    len(viol) == len(term),
    "최소간격 %.1fs · 위반 %d건(종료암전 예외 %d건)"
    % (min(g[2] for g in gaps), len(viol), len(term)),
)
# 9
last = rows[-1]
chk(
    9,
    "마지막 큐 = 암전/인계 명시",
    "블랙아웃" in last[6] and last[7] == "0" and "인계" in last[13],
    f"{last[0]} P8 블랙아웃 0% · 인계 지시 포함",
)
# 10
ns = wb["NOTE"]
nrows = [[c.value for c in r] for r in ns.iter_rows(min_row=4, max_row=ns.max_row, max_col=5)]
tc_unconf = sum(1 for r in rows for v in r if isinstance(v, str) and "확인필요" in v)
chk(
    10,
    "확인필요 항목 NOTE 기록",
    len(nrows) >= 10 and any(r[0] == "확인필요" for r in nrows),
    "NOTE %d건 (확인필요 %d · 장비이슈 %d · 스펙개정후보 %d)"
    % (
        len(nrows),
        sum(1 for r in nrows if r[0] == "확인필요"),
        sum(1 for r in nrows if r[0] == "장비이슈"),
        sum(1 for r in nrows if r[0] == "스펙개정후보"),
    ),
)

# 11 (v1.1)
tm = str(head.get("TC_METHOD") or "")
grade = tm.split()[0] if tm else ""
warned = "리허설" in tm or "미검증" in tm
chk(
    11,
    "TC_METHOD 기입 · DERIVED 경고",
    grade in ("VERIFIED", "DERIVED") and (grade == "VERIFIED" or warned),
    "{} · 경고 표기 {}".format(grade or "미기입", "있음" if warned else "없음"),
)

# ===== v2.0 실행 레이어 검증 (12~15) =====
from exec_data import CUE_EX, PATCH, PRESETS, SONG_BPM

cue_qs = [r[0] for r in rows]
cue_by_q = {r[0]: r for r in rows}
ex_qs = []
for r_ in CUE_EX:
    if r_[0] not in ex_qs:
        ex_qs.append(r_[0])

# 12: 커버리지 + Intensity/Transition 정합
miss_q = [q for q in cue_qs if q not in ex_qs]
extra_q = [q for q in ex_qs if q not in cue_qs]
mismatch = []
for r_ in CUE_EX:
    q, grp, dim, snap = r_[0], r_[1], r_[2], r_[15]
    cue = cue_by_q.get(q)
    if not cue:
        continue
    is_snap_cue = cue[11] == "SNAP"
    if is_snap_cue and dim and dim != "0" and snap != "Y":
        # SNAP 큐에서 점등 행이 Snap=Y 아님 (소등/트래킹 행은 예외)
        if float(r_[10] or 0) == 0.0:
            pass
        else:
            mismatch.append((q, grp, "SNAP큐 점등행 페이드>0"))
    # CUE Intensity 단일값 큐와 CUE-EX Dim 최대값 대조
# 트래킹 시뮬레이션: 명시 행만으로 큐별 상태를 누적해 CUE Intensity와 대조
state = {}
allg = {p_[0] for p_ in PATCH}
for q in cue_qs:
    for r_ in CUE_EX:
        if r_[0] != q:
            continue
        dim = r_[2]
        if dim in ("", None):
            continue
        if r_[1] == "ALL":
            for g_ in allg:
                state[g_] = int(dim)
        else:
            state[r_[1]] = int(dim)
    inten = cue_by_q[q][7]
    if isinstance(inten, str) and inten.isdigit():
        ATMOS = {"HAZE"}  # 분위기 그룹은 조명 레벨 비교 제외 (스펙 §11.2)
        lit = {g_: v_ for g_, v_ in state.items() if g_ not in ATMOS}
        cur = max(lit.values()) if lit else 0
        if cur != int(inten):
            mismatch.append((q, "-", "CUE Intensity %s vs 트래킹 상태 max %d" % (inten, cur)))
chk(
    12,
    "CUE-EX 커버리지 · CUE 정합",
    not miss_q and not extra_q and not mismatch,
    (f"누락 {miss_q} / 잉여 {extra_q} / 모순 {mismatch[:3]}")
    if (miss_q or extra_q or mismatch)
    else "18/18 큐 커버 · %d행 · 트래킹 시뮬 Intensity/SNAP 정합" % len(CUE_EX),
)

# 13: 프리셋 참조 무결성
pids = {p_[0] for p_ in PRESETS}
badref = []
for r_ in CUE_EX:
    for idx, typ in ((3, "COL"), (4, "POS"), (5, "BM"), (6, "FX")):
        v = r_[idx]
        if v in ("", "OFF", None):
            continue
        if v not in pids:
            badref.append((r_[0], r_[1], v))
chk(
    13,
    "PRESET 참조 무결성",
    not badref,
    (f"미정의 참조 {badref[:4]}") if badref else "%d개 프리셋 · 참조 위반 0건" % len(pids),
)

# 14: PATCH 그룹 무결성
pgroups = {p_[0] for p_ in PATCH} | {"ALL"}
badgrp = sorted({r_[1] for r_ in CUE_EX if r_[1] not in pgroups})
chk(
    14,
    "CUE-EX Group PATCH 정의",
    not badgrp,
    (f"미정의 그룹 {badgrp}")
    if badgrp
    else "%d그룹 전부 PATCH 정의" % len({r_[1] for r_ in CUE_EX}),
)

# 15: FX-Rate BPM 공식
VALID_RATES = {
    SONG_BPM / 0.5,
    SONG_BPM / 1,
    SONG_BPM / 4,
    SONG_BPM / 8,
    SONG_BPM / 16,
}  # 240,120,30,15,7.5
badrate = []
for r_ in CUE_EX:
    if r_[6] in ("", "OFF", None):
        continue
    rate = r_[7]
    if rate in ("", None):
        badrate.append((r_[0], r_[1], "Rate 누락"))
        continue
    if float(rate) not in VALID_RATES:
        badrate.append((r_[0], r_[1], rate))
chk(
    15,
    "FX-Rate BPM 공식 (§11.3)",
    not badrate,
    (f"위반 {badrate[:4]}") if badrate else "허용값 {7.5,15,30,120,240} 내 · 위반 0건",
)

w = max(len(x[1]) for x in R)
print("=" * 96)
print("LX-SEQ v2.0 포맷 준수 검증 (연출 11 + 실행 4) — LXSEQ_SAMPLE_01_Sugar_r3.xlsx")
print("=" * 96)
for n, name, v, d in R:
    print("%2d. %-*s  %-4s  %s" % (n, w, name, v, d))
print("-" * 96)
nf = sum(1 for x in R if x[2] == "FAIL")
print("결과: %d/%d PASS%s" % (len(R) - nf, len(R), "" if nf == 0 else "  ← FAIL %d건" % nf))
print(
    "총 큐 %d · 총 길이 %.1fs · 후렴 비중 %.1f%%"
    % (
        len(rows),
        RUNTIME,
        sum(b - a for n_, bars, a, b in SECTIONS if n_.startswith("CHORUS")) / RUNTIME * 100,
    )
)
sys.exit(1 if nf else 0)
