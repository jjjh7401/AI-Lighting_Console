"""실행 레이어 시트(PATCH/PRESET/CUE-EX) 추가 + CUE-EX CSV 내보내기 (r2)"""

import csv
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

# 입출력 위치는 이 스크립트 위치에서 유도한다. 기계마다 다른 절대경로를
# 박아두면 그 기계 밖에서는 돌지 않는다.
SONG_DIR = os.environ.get("LXSEQ_SONG_OUT") or os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "03_곡파일_Sugar"
)
os.makedirs(SONG_DIR, exist_ok=True)
from exec_data import CUE_EX, EX_HEADERS, EXEC_NOTES, PATCH, PRESETS
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SRC = os.path.join(SONG_DIR, "LXSEQ_SAMPLE_01_Sugar_r3.xlsx")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR_FILL = PatternFill("solid", fgColor="1F4E3D")  # 실행 레이어는 녹색 계열로 구분
HDR_FONT = Font(name="맑은 고딕", size=9, bold=True, color="FFFFFF")
BODY = Font(name="맑은 고딕", size=9)
MONO = Font(name="Consolas", size=9)

wb = load_workbook(SRC)


def header(ws, row, cols):
    for i, h in enumerate(cols, start=1):
        c = ws.cell(row, i, h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------- PATCH ----------
ps = wb.create_sheet("PATCH")
ps.cell(1, 1, "PATCH — 장비 패치 (기종·주소 전부 가정값, 현장 확정 필요)").font = Font(
    name="맑은 고딕", size=11, bold=True, color="1F4E3D"
)
header(ps, 3, ["Group", "Fixture Type", "Mode", "Count", "Fixture ID", "Universe.Address", "Note"])
r = 4
for row in PATCH:
    for i, v in enumerate(row, start=1):
        c = ps.cell(r, i, v)
        c.font = MONO if i in (1, 5, 6) else BODY
        c.border = BORDER
        if str(v) == "확인필요":
            c.font = Font(name="맑은 고딕", size=9, bold=True, color="B03030")
    r += 1
for col, w in zip("ABCDEFG", [10, 24, 8, 7, 15, 16, 30]):
    ps.column_dimensions[col].width = w

# ---------- PRESET ----------
qs = wb.create_sheet("PRESET")
qs.cell(
    1, 1, "PRESET — 프리셋 정의 (POS는 현장 레코드 대상 · COL은 연출 팔레트 P.xx와 1:1)"
).font = Font(name="맑은 고딕", size=11, bold=True, color="1F4E3D")
header(qs, 3, ["Preset ID", "이름(무대 의미)", "대상 Group", "절대값", "MA3 Pool", "Note"])
TYPE_FILL = {"POS": "E2ECFF", "COL": "FFE2F0", "BM": "FFF2DC", "FX": "E4F8F8"}
r = 4
for row in PRESETS:
    fill = PatternFill("solid", fgColor=TYPE_FILL.get(row[0].split(".")[0], "FFFFFF"))
    for i, v in enumerate(row, start=1):
        c = qs.cell(r, i, v)
        c.font = MONO if i == 1 else BODY
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=(i == 4))
        if i == 1:
            c.fill = fill
        if str(v) == "현장 레코드":
            c.font = Font(name="맑은 고딕", size=9, bold=True, color="B03030")
    r += 1
for col, w in zip("ABCDEF", [10, 28, 18, 32, 10, 24]):
    qs.column_dimensions[col].width = w

# ---------- CUE-EX ----------
xs = wb.create_sheet("CUE-EX")
xs.cell(
    1, 1, "CUE-EX — 큐 × 그룹별 파라미터·타이밍 (long format · 빈칸=트래킹 · 타깃 grandMA3)"
).font = Font(name="맑은 고딕", size=11, bold=True, color="1F4E3D")
xs.cell(
    2,
    1,
    "FX-Rate는 BPM 단위, 곡 120BPM 기준: @1/8=240 · @1beat=120 · @1bar=30 · @2bar=15 · @4bar=7.5  |  Speed Master 120BPM 종속 권장",
).font = Font(name="맑은 고딕", size=9, italic=True, color="4A5568")
HR = 4
header(xs, HR, EX_HEADERS)
Q_FILL = {}
palette_cycle = ["F5F7FA", "EDF1F6"]
qseen = []
r = HR + 1
for row in CUE_EX:
    q = row[0]
    if q not in qseen:
        qseen.append(q)
    fill = PatternFill("solid", fgColor=palette_cycle[qseen.index(q) % 2])
    for i, v in enumerate(row, start=1):
        c = xs.cell(r, i, v)
        c.font = MONO if i not in (17,) else BODY
        c.border = BORDER
        c.fill = fill
        c.alignment = Alignment(
            vertical="center", horizontal="center" if i < 17 else "left", wrap_text=(i == 17)
        )
        if i == 16 and v == "Y":
            c.font = Font(name="Consolas", size=9, bold=True, color="B03030")
    r += 1
widths = [7, 9, 6, 8, 8, 7, 7, 8, 9, 8, 7, 7, 7, 7, 7, 5, 42]
for idx, w in enumerate(widths):
    xs.column_dimensions[chr(65 + idx) if idx < 26 else "A"].width = w
xs.column_dimensions["Q"].width = 42
xs.freeze_panes = "C5"
xs.page_setup.orientation = "landscape"
xs.page_setup.fitToWidth = 1
xs.sheet_properties.pageSetUpPr.fitToPage = True
xs.print_title_rows = "%d:%d" % (HR, HR)

# ---------- NOTE 시트에 실행 레이어 노트 추가 ----------
ns = wb["NOTE"]
GUBUN_COLOR = {
    "확인필요": "FFE0E0",
    "장비이슈": "FFF0D8",
    "스펙개정후보": "E2ECFF",
    "리허설변경": "E8F5E0",
}
start = ns.max_row + 1
for g, q, txt, d, st in EXEC_NOTES:
    row = [g, q, txt, d, st]
    for i, v in enumerate(row, start=1):
        c = ns.cell(start, i, v)
        c.font = BODY
        c.border = BORDER
        c.alignment = Alignment(
            vertical="center",
            wrap_text=(i == 3),
            horizontal="center" if i in (1, 2, 4, 5) else "left",
        )
        if i == 1:
            c.fill = PatternFill("solid", fgColor=GUBUN_COLOR.get(g, "FFFFFF"))
            c.font = Font(name="맑은 고딕", size=9, bold=True)
    start += 1

wb.save(SRC)
print("exec sheets added:", wb.sheetnames)

# ---------- CUE-EX CSV (기계 정본) ----------
csv_path = os.path.join(SONG_DIR, "LXSEQ_SAMPLE_01_Sugar_r3.cue-ex.csv")
with open(csv_path, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(EX_HEADERS)
    for row in CUE_EX:
        w.writerow(row)
print("csv:", csv_path, len(CUE_EX), "rows")
